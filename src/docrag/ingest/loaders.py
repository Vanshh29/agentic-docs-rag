"""Loaders for PDF (per page) and DOCX (per heading section)."""

from __future__ import annotations

import logging
from pathlib import Path

from langchain_core.documents import Document

log = logging.getLogger(__name__)

SUPPORTED = {".pdf", ".docx"}


def load_pdf(path: Path, source_file: str | None = None) -> list[Document]:
    from pypdf import PdfReader

    name = source_file or path.name
    docs: list[Document] = []
    for i, page in enumerate(PdfReader(str(path)).pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            docs.append(Document(page_content=text, metadata={"source_file": name, "page": i}))
    return docs


def load_docx(path: Path, source_file: str | None = None) -> list[Document]:
    """Split on Heading styles so each Document is one section."""
    from docx import Document as DocxDocument

    name = source_file or path.name
    sections: list[tuple[str, list[str]]] = [("Introduction", [])]
    for p in DocxDocument(str(path)).paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name if p.style is not None and p.style.name else "") or ""
        if style.startswith(("Heading", "Title", "Überschrift")):
            sections.append((text, []))
        else:
            sections[-1][1].append(text)
    return [
        Document(
            page_content=f"{title}\n" + "\n".join(body),
            metadata={"source_file": name, "section": title},
        )
        for title, body in sections
        if body
    ]


def load_path(path: Path) -> tuple[list[Document], list[Path], list[tuple[str, str]]]:
    """Load all supported files under `path`, isolating failures per file.

    `source_file` metadata is the POSIX path relative to the ingest root (the directory itself,
    or the parent of a single file), so same-named files in different folders stay distinct.

    Returns (documents, skipped_files, failed) where failed holds (source_file, error message).
    """
    root = path.parent if path.is_file() else path
    files = [path] if path.is_file() else sorted(p for p in path.rglob("*") if p.is_file())
    docs: list[Document] = []
    skipped: list[Path] = []
    failed: list[tuple[str, str]] = []
    for f in files:
        rel = f.relative_to(root).as_posix()
        ext = f.suffix.lower()
        if f.name.startswith(".") or f.name == "README.md" or ext not in SUPPORTED:
            log.warning("Skipping unsupported file: %s", f)
            skipped.append(f)
            continue
        try:
            docs.extend(load_pdf(f, rel) if ext == ".pdf" else load_docx(f, rel))
        except Exception as e:  # corrupt/encrypted/empty files must not abort the batch
            log.warning("Failed to load %s: %s: %s", f, type(e).__name__, e)
            failed.append((rel, f"{type(e).__name__}: {e}"))
    return docs, skipped, failed
