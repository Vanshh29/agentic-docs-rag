"""Builders for small PDF/DOCX files used in ingestion tests."""

from __future__ import annotations

from pathlib import Path


def make_pdf(path: Path, pages: list[str]) -> Path:
    """Write a PDF with one page per string; an empty string yields a blank page."""
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path))
    for text in pages:
        y = 800
        for line in text.splitlines():
            c.drawString(50, y, line)
            y -= 14
        c.showPage()
    c.save()
    return path


def make_docx(path: Path, sections: dict[str, str]) -> Path:
    from docx import Document

    d = Document()
    for title, body in sections.items():
        d.add_heading(title, level=1)
        d.add_paragraph(body)
    d.save(str(path))
    return path
