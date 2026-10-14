"""Generate the synthetic sample manuals in data/samples (original, fictional content).

Usage: python scripts/make_samples.py
"""

from __future__ import annotations

from pathlib import Path

from docx import Document as Docx
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

OUT = Path(__file__).resolve().parents[1] / "data" / "samples"

Section = tuple[str, list[str]]

ESPRESSO_EN: list[Section] = [
    (
        "1. Safety Instructions",
        [
            "Read this manual before using the AromaPro E200 espresso machine. Keep it for "
            "future reference.",
            "Operate the machine only on a stable, level, heat-resistant surface. The machine "
            "is rated for 230 V, 50 Hz and consumes 1450 W. Never immerse the machine, the "
            "power cord or the plug in water or any other liquid.",
            "The appliance may be used by children from 8 years of age only under supervision. "
            "Keep the machine and its cord out of reach of children younger than 8 years.",
            "The ambient temperature during operation must not exceed 35 degrees Celsius. "
            "The steam wand and the brew group become very hot; do not touch them during use.",
        ],
    ),
    (
        "2. Daily Operation",
        [
            "Filling the water tank: remove the tank at the rear and fill it with fresh, cold "
            "water up to the MAX mark. The tank capacity is 1.8 litres.",
            "Heat-up: after switching on, the machine needs approximately 25 seconds to heat up. "
            "The ready light turns steady when brewing is possible.",
            "Brewing: press the single-shot button for 30 ml or the double-shot button for 60 ml. "
            "The pump pressure is 15 bar. The portafilter holds 8 g of ground coffee for a "
            "single shot and 16 g for a double shot.",
            "Frothing milk: press the steam button and wait about 40 seconds until the steam "
            "light stays on. Purge the steam wand briefly before and after every use and wipe "
            "it with a damp cloth immediately after frothing.",
        ],
    ),
    (
        "3. Cleaning and Maintenance",
        [
            "Daily: empty the drip tray and rinse the portafilter and filter baskets with warm "
            "water.",
            "Weekly: backflush the brew group. Insert the blind filter basket with one cleaning "
            "tablet and run 5 cycles of 10 seconds each, then rinse with clean water.",
            "Descaling: descale the machine every 3 months, or immediately when the descale "
            "light blinks. Mix 1 litre of descaling solution at a ratio of 1:3 with water. "
            "The descaling programme takes about 20 minutes. Rinse the tank and run two full "
            "tanks of fresh water through the machine afterwards.",
            "Water filter: replace the water filter cartridge every 2 months or after 60 litres, "
            "whichever comes first. Grinder: clean the burrs monthly with the supplied brush.",
        ],
    ),
    (
        "4. Troubleshooting and Warranty",
        [
            "Error E01: the water tank is empty or not seated correctly. Refill the tank and "
            "push it firmly into place.",
            "Error E02: overheating protection (thermal fuse) triggered. Unplug the machine and "
            "let it cool down for 30 minutes before restarting.",
            "Error E04: pump blockage, usually caused by limescale. Run the descaling programme. "
            "Error E07: the grinder is jammed. Switch the machine off, remove the beans from "
            "the hopper and clean the burrs.",
            "Coffee tastes too weak: select a finer grind. The grinder has settings 1 to 12, "
            "where 1 is the finest; the factory default is 5.",
            "Warranty: 24 months from the date of purchase. Contact customer service on "
            "+49 800 555 0101 (fictional number).",
        ],
    ),
]

CHARGER_DE: list[Section] = [
    (
        "1. Sicherheitshinweise",
        [
            "Das Staplerladegerät VoltMax LX-48 darf nur von geschultem Personal bedient werden.",
            "Der Ladeplatz muss gut belüftet sein, da beim Laden von Blei-Säure-Batterien "
            "explosives Wasserstoffgas entstehen kann. Rauchen, offenes Feuer und Funkenbildung "
            "sind im Ladebereich verboten.",
            "Beim Arbeiten an der Batterie sind Schutzbrille und säurefeste Handschuhe zu tragen.",
            "Netzanschluss: 400 V dreiphasig, abgesichert mit 16 A. Schutzart IP21. Die "
            "zulässige Umgebungstemperatur beträgt 0 bis 40 °C.",
        ],
    ),
    (
        "2. Inbetriebnahme und Laden",
        [
            "Das Ladegerät ist für Batterien mit 48 V Nennspannung und einer Kapazität von "
            "200 bis 600 Ah ausgelegt. Der maximale Ladestrom beträgt 80 A.",
            "Verbinden Sie zuerst den Batteriestecker mit dem Ladegerät. Das Gerät erkennt die "
            "Batteriespannung automatisch. Drücken Sie anschließend die Start-Taste für "
            "2 Sekunden.",
            "Ladeprogramme: P1 Blei-Säure (Standard), P2 Gel-Batterie, P3 Lithium-Ionen. "
            "Wählen Sie das Programm mit der Programm-Taste.",
            "Die Ladezeit beträgt etwa 8 Stunden für eine Batterie mit 500 Ah bei einer "
            "Entladungstiefe von 80 %. Die grüne LED blinkt während des Ladens und leuchtet "
            "dauerhaft, wenn die Ladung abgeschlossen ist.",
            "Eine Ausgleichsladung sollte einmal pro Woche durchgeführt werden und dauert "
            "zusätzlich etwa 3 Stunden.",
        ],
    ),
    (
        "3. Wartung und Pflege",
        [
            "Monatlich: Kontakte und Stecker prüfen und reinigen sowie alle Kabel auf "
            "Beschädigungen kontrollieren.",
            "Der Lüfterfilter muss alle 6 Monate gewechselt werden, bei staubiger Umgebung früher.",
            "Eine jährliche Prüfung durch eine Elektrofachkraft nach DGUV Vorschrift 3 ist "
            "verpflichtend.",
            "Reinigen Sie das Gehäuse nur mit einem trockenen oder leicht feuchten Tuch. "
            "Die Verwendung eines Hochdruckreinigers ist nicht zulässig.",
            "Batteriewasser (destilliertes Wasser) darf nur nach dem Laden nachgefüllt werden, "
            "niemals davor.",
        ],
    ),
    (
        "4. Fehlerbehebung und Garantie",
        [
            "Fehler F1: Überspannung im Netz. Prüfen Sie die Netzspannung.",
            "Fehler F2: Batterie nicht erkannt. Prüfen Sie den Batteriestecker und die Kontakte.",
            "Fehler F3: Übertemperatur. Prüfen Sie Lüfter und Lüfterfilter und lassen Sie das "
            "Gerät 20 Minuten abkühlen.",
            "Fehler F5: Ladezeitüberschreitung nach 14 Stunden. Lassen Sie die Batterie "
            "prüfen, möglicherweise ist eine Zelle defekt. Fehler F7: Kommunikationsfehler "
            "mit dem Batteriemanagementsystem (nur Lithium-Ionen).",
            "Eine dauerhaft leuchtende rote LED zeigt eine Störung an.",
            "Garantie: 36 Monate ab Kaufdatum.",
        ],
    ),
]

PURIFIER_DE: list[Section] = [
    (
        "Sicherheit",
        [
            "Der Luftreiniger ClearAir P5 ist nur für den Gebrauch in Innenräumen bestimmt. "
            "Stellen Sie das Gerät mindestens 30 cm von Wänden entfernt auf.",
            "Betreiben Sie das Gerät nicht in Räumen mit hoher Feuchtigkeit wie Badezimmern. "
            "Ziehen Sie vor jeder Reinigung den Netzstecker.",
        ],
    ),
    (
        "Technische Daten",
        [
            "Geeignet für Räume bis 45 m². Der Luftdurchsatz (CADR) beträgt 320 m³/h. "
            "Die Leistungsaufnahme beträgt 55 W.",
            "Die Lautstärke beträgt 24 dB im Schlafmodus und 52 dB auf Stufe 3.",
        ],
    ),
    (
        "Betriebsarten",
        [
            "Automatikmodus: Der Partikelsensor regelt die Lüfterstufe selbstständig nach der "
            "Luftqualität.",
            "Schlafmodus: leiser Betrieb, die Anzeige wird abgedunkelt. Turbomodus: maximale "
            "Leistung für 30 Minuten, danach wechselt das Gerät in den Automatikmodus.",
            "Über die ClearAir App kann das Gerät per WLAN gesteuert werden. Es werden nur "
            "2,4-GHz-Netzwerke unterstützt.",
        ],
    ),
    (
        "Filterwechsel",
        [
            "Der HEPA-Filter H13 muss alle 12 Monate gewechselt werden. Der Aktivkohlefilter "
            "muss alle 6 Monate gewechselt werden.",
            "Den Vorfilter alle 2 Wochen mit einem Staubsauger absaugen.",
            "Die Filterwechselanzeige leuchtet rot nach 8000 Betriebsstunden. Nach dem Wechsel "
            "halten Sie die Filter-Taste 5 Sekunden gedrückt, um die Anzeige zurückzusetzen.",
        ],
    ),
    (
        "Störungen und Garantie",
        [
            "E1: Lüfter blockiert. Schalten Sie das Gerät aus und entfernen Sie Fremdkörper.",
            "E2: Sensor verschmutzt. Reinigen Sie die Sensoröffnung vorsichtig mit einem "
            "trockenen Wattestäbchen.",
            "E3: Filterklappe nicht geschlossen. Schließen Sie die Klappe, bis sie einrastet.",
            "Garantie: 24 Monate ab Kaufdatum.",
        ],
    ),
]


def make_pdf(path: Path, title: str, sections: list[Section]) -> None:
    styles = getSampleStyleSheet()
    story: list[object] = [Paragraph(title, styles["Title"]), Spacer(1, 12)]
    for i, (heading, paras) in enumerate(sections):
        if i:
            story.append(PageBreak())
        story.append(Paragraph(heading, styles["Heading2"]))
        for p in paras:
            story += [Paragraph(p, styles["BodyText"]), Spacer(1, 8)]
    SimpleDocTemplate(str(path), pagesize=A4, title=title).build(story)


def make_docx(path: Path, title: str, sections: list[Section]) -> None:
    doc = Docx()
    doc.add_heading(title, level=0)
    for heading, paras in sections:
        doc.add_heading(heading, level=1)
        for p in paras:
            doc.add_paragraph(p)
    doc.save(str(path))


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    make_pdf(OUT / "espresso_machine_manual_en.pdf", "AromaPro E200 Espresso Machine", ESPRESSO_EN)
    make_pdf(OUT / "forklift_charger_manual_de.pdf", "VoltMax LX-48 Staplerladegerät", CHARGER_DE)
    make_docx(OUT / "air_purifier_manual_de.docx", "ClearAir P5 Luftreiniger", PURIFIER_DE)
    print("Wrote samples to", OUT)


if __name__ == "__main__":
    main()
